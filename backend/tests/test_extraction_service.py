from io import BytesIO

import docx
import pytest

from app.services.extraction_service import ExtractionError, extract_text


def test_extracts_plain_text():
    assert extract_text("note.txt", b"Hello invoice") == "Hello invoice"


def test_extracts_markdown():
    assert extract_text("readme.md", b"# Title") == "# Title"


def test_extracts_docx():
    document = docx.Document()
    document.add_paragraph("Quarterly report")
    document.add_paragraph("Revenue grew by ten percent.")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_text("report.docx", buffer.getvalue())
    assert "Quarterly report" in text
    assert "Revenue grew by ten percent." in text


def test_rejects_empty_file():
    with pytest.raises(ExtractionError):
        extract_text("empty.txt", b"   ")


def test_rejects_unsupported_extension():
    with pytest.raises(ExtractionError):
        extract_text("archive.zip", b"PK\x03\x04")


def test_rejects_malformed_pdf():
    with pytest.raises(ExtractionError):
        extract_text("broken.pdf", b"not really a pdf")
